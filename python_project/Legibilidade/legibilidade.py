#Biblioteca geral, herdada do código anterior
import re
#Biblioteca utilizada para contagem de sentenças
import nltk
nltk.download('punkt')
#Biblioteca para criação de janelas para acesso dos caminhos necessários
from tkinter import filedialog
#biblioteca para elaboração de planilha de resultados
from openpyxl import Workbook
#bibliotecas para extração de textos, herdada do código Final.py
from io import StringIO
from pdfminer.converter import TextConverter
from pdfminer.layout import LAParams
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.pdfpage import PDFPage
from pdfminer.pdfparser import PDFParser
#biblioteca para acesso e movimentação entre paths
import os
#biblioteca para contagem de sílabas
import pyphen


'''
1 - Quantidade de letras.
2 - Quantidade de sílabas. 
3 - Quantidade de palavras.
4 - Quantidade de palavras complexas.
Palavras complexas = palavras com 3 ou mais sílabas.
5 - Quantidade de sentenças.
Sentença = oração (é toda frase que contém um ou mais verbos), normalmente dentro de um parágrafo, a sentença ou oração, é delimitada entre o início da frase (letra maiúscula) até o sinal de pontuação.
Exemplo de uma sentença: 
As demonstrações contábeis foram elaboradas seguindo os princípios contábeis. 
6 – Quantidade de páginas.
'''

hifenizador = pyphen.Pyphen(lang='pt')

janela_data = filedialog.askdirectory()
lista_relatorio = os.listdir(janela_data)
lista_relatorio.sort()
#print(lista_relatorio)
janela_saida = filedialog.askdirectory()

planilha1 = Workbook()
data_geral = planilha1.active
data_geral.title = "Data Geral"
data_geral["A2"] = "Empresas"
data_geral["B1"] = "Qt. de letras"
data_geral["C1"] = "Qt. de sílabas"
data_geral["D1"] = "Qt. de palavras"
data_geral["E1"] = "Qt. de pal. complexas"
data_geral["F1"] = "Qt. de Sentenças"
data_geral["G1"] = "Qt. de páginas"

planilha2 = Workbook()
data_local = planilha2.active
data_local.title = "Data Local"
data_local["A1"] = "Empresas"
data_local["B1"] = "Qt. de letras"
data_local["C1"] = "Qt. de sílabas"
data_local["D1"] = "Qt. de palavras"
data_local["E1"] = "Qt. de pal. complexas"
data_local["F1"] = "Qt. de Sentenças"
data_local["G1"] = "Qt. de páginas"
data_local["H1"] = "CFRE"
data_local["I1"] = "CFKGL"
indice_local = 2

def extract_text(path):
    '''
    Extração de textos dos arquivos em pdf
    '''
    output_string = StringIO()
    with open(path,'rb') as in_file:
        parser = PDFParser(in_file)
        doc = PDFDocument(parser)
        rsrcmgr = PDFResourceManager()
        device = TextConverter(rsrcmgr, output_string, laparams=LAParams())
        interpreter = PDFPageInterpreter(rsrcmgr,device)
        for page in PDFPage.create_pages(doc):
            interpreter.process_page(page)
        #variavel de retorno da extração do texto
        text = output_string.getvalue()
        bad = ['\n', '\r','\t','\n\x0c','1','2','3','4','5','6','7','8','9','0','*','-','/','%','$','(.)','(..)','(...)','()','( )','(   )',',.','..','+']
        for name in bad:
            text = text.replace(name, '')
        return text

def count_letters(text):
    """
    Retorna o número de letras em um texto.
    """
    text = re.sub(r'[^a-zA-Z]', '', text)
    return len(text)

def count_syllables(word):
    """
    Retorna o número de sílabas em uma palavra.
    """
    word = word.lower().strip(".:;?!,")

    # Usa o hifenizador para dividir a palavra em sílabas
    syllables = hifenizador.inserted(word)
    count = syllables.count('-') + 1

    return count

def count_words(text):
    """
    Retorna o número de palavras em um texto.
    """
    words = re.findall(r'\b\w+\b', text)
    return len(words)


def count_complex_words(text):
    """
    Retorna o número de palavras complexas em um texto.
    """
    words = re.findall(r'\b\w+\b', text)
    count = 0
    for word in words:
        num_syllables = count_syllables(word)
        if num_syllables >= 3:
            count += 1
    return count


def count_sentences(text):
    """
    Retorna o número de sentenças em um texto.
    """
    sentences = nltk.sent_tokenize(text)
    return len(sentences)       

def count_pages(path):
    """
    Retorna o número de páginas de um arquivo PDF.
    """
    with open(path, 'rb') as file:
        parser = PDFParser(file)
        document = PDFDocument(parser)
        return len(list(PDFPage.create_pages(document)))

def calculate_flesch_reading_ease(text):
    """
    Retorna o índice de legibilidade de Flesch Reading Ease. Na planilha ele está indicado na coluna cfre.
    """
    num_sentences = count_sentences(text)
    num_words = count_words(text)
    num_syllables = sum([count_syllables(word) for word in text.split()])
    FRE = 206.835 - 1.015 * (num_words / num_sentences) - 84.6 * (num_syllables / num_words)
    return round(FRE, 2)

def calculate_flesch_kincaid_grade_level(text):
    """
    Retorna o índice de legibilidade de Flesch-Kincaid Grade Level. Na planilha ele está indicado na coluna cfkgl.
    """
    num_sentences = count_sentences(text)
    num_words = count_words(text)
    num_syllables = sum([count_syllables(word) for word in text.split()])
    FKGL = 0.39 * (num_words / num_sentences) + 11.8 * (num_syllables / num_words) - 15.59
    return round(FKGL, 2)

def clean_text(text):
    """
    Remove caracteres especiais e transforma o texto em letras minúsculas.
    """
    text = re.sub(r'[^\w\s]', '', text)
    text = text.lower()
    return text

t = len(lista_relatorio)
#print(t)
i = 0
while i < t:
    print(25*"*-")
    print(f'Estamos no arquivo {i + 1}, no total temos {t}')
    caminho_relatorio = janela_data + str('/') + lista_relatorio[i]
    text = extract_text(caminho_relatorio)
    #text = clean_text(text)
    #print(text)
    contar_letras = count_letters(text)
    contar_silabas = sum([count_syllables(word) for word in text.split()])
    contar_palavras = count_words(text)
    contar_palavras_c = count_complex_words(text)
    contar_sentenca = count_sentences(text)
    contar_paginas = count_pages(caminho_relatorio)
    # Calcula os índices de legibilidade
    FRE = calculate_flesch_reading_ease(text)
    FKGL = calculate_flesch_kincaid_grade_level(text)
    print(lista_relatorio[i])
    print([contar_letras, contar_silabas, contar_palavras, contar_palavras_c, contar_sentenca, contar_paginas, FRE, FKGL])    
    
    #print(lista_relatorio[i])
    name = lista_relatorio[i].split('_')
    #print(name[0])
    indice_planilha_A = 'A' + str(indice_local)
    indice_planilha_B = 'B' + str(indice_local)
    indice_planilha_C = 'C' + str(indice_local)
    indice_planilha_D = 'D' + str(indice_local)
    indice_planilha_E = 'E' + str(indice_local)
    indice_planilha_F = 'F' + str(indice_local)
    indice_planilha_G = 'G' + str(indice_local)
    indice_planilha_H = 'H' + str(indice_local)
    indice_planilha_I = 'I' + str(indice_local)
    
    data_local[indice_planilha_A] = name[0]
    data_local[indice_planilha_B] = contar_letras
    data_local[indice_planilha_C] = contar_silabas
    data_local[indice_planilha_D] = contar_palavras
    data_local[indice_planilha_E] = contar_palavras_c
    data_local[indice_planilha_F] = contar_sentenca
    data_local[indice_planilha_G] = contar_paginas
    data_local[indice_planilha_H] = FRE
    data_local[indice_planilha_I] = FKGL
    
    
    
    i += 1
    indice_local += 1

print("Análise finalizada")
name1 = '/Analise_Legibilidade.xlsx'
path_saida = janela_data + name1
planilha2.save(path_saida)